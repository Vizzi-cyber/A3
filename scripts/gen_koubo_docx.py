"""从「演示视频口播稿.md」生成同名 Word 文档。

用法：
    python scripts/gen_koubo_docx.py

格式约定（见 md 首段「文档格式说明」，与既有交付稿一致）：
  - A4，上下边距 2.5cm，左右 2cm，1.5 倍行距
  - 一级标题：黑体二号居中；二级：黑体小三；三级：黑体四号
  - 正文：宋体小四
  - 「画面操作」整段：灰色字体，与口播/旁白区分
  - 正文中的 **加粗** 标记会转成 Word 加粗
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "文档信息" / "演示视频口播稿.md"
DOCX_PATH = ROOT / "文档信息" / "演示视频口播稿.docx"

GRAY = RGBColor(0x80, 0x80, 0x80)
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _style_font(run, cn_font: str, size_pt: float, *, bold: bool = False,
                color: RGBColor | None = None) -> None:
    """设置 run 的中英文字体、字号、加粗与颜色。"""
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = cn_font
    # 中文字体需额外写 w:eastAsia，否则 python-docx 只设了西文字体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
    if color is not None:
        run.font.color.rgb = color


def _add_body(doc: Document, text: str, *, gray: bool = False) -> None:
    """添加正文段落，支持 **加粗** 内联标记。"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)

    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            _style_font(p.add_run(text[pos:m.start()]), "宋体", 12,
                        color=GRAY if gray else None)
        _style_font(p.add_run(m.group(1)), "宋体", 12, bold=True,
                    color=GRAY if gray else None)
        pos = m.end()
    if pos < len(text):
        _style_font(p.add_run(text[pos:]), "宋体", 12,
                    color=GRAY if gray else None)


def _add_heading(doc: Document, text: str, level: int) -> None:
    """添加标题：一级黑体二号居中，二级黑体小三，三级黑体四号。"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        _style_font(p.add_run(text), "黑体", 22, bold=True)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        _style_font(p.add_run(text), "黑体", 15, bold=True)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        _style_font(p.add_run(text), "黑体", 14, bold=True)


def main() -> int:
    if not MD_PATH.exists():
        print(f"[错误] 找不到源文件：{MD_PATH}", file=sys.stderr)
        return 1

    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    for raw in MD_PATH.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.strip() == "---":
            continue
        if line.startswith("### "):
            _add_heading(doc, line[4:].strip(), 3)
        elif line.startswith("## "):
            _add_heading(doc, line[3:].strip(), 2)
        elif line.startswith("# "):
            _add_heading(doc, line[2:].strip(), 1)
        elif line.lstrip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.5
            _style_font(p.add_run(line.lstrip()[2:].strip()), "宋体", 12)
        else:
            _add_body(doc, line.strip(), gray=line.startswith("画面操作"))

    doc.save(DOCX_PATH)
    print(f"[OK] 已生成：{DOCX_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
